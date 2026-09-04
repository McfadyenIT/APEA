"""Reporter — JMeter-style HTML dashboard + Excel workbook.

Consumes the analyzer output and writes:
    reports/performance_report.html  (dark-theme interactive dashboard)
    reports/performance_report.xlsx  (multi-sheet workbook)
"""
from __future__ import annotations

import csv
import html
import json
from pathlib import Path


def _read_history(run_dir: Path) -> list[dict]:
    path = run_dir / "results" / "locust_stats_history.csv"
    if not path.exists():
        return []
    try:
        with open(path, newline="", encoding="utf-8") as fh:
            rows = [r for r in csv.DictReader(fh) if r.get("Name") == "Aggregated"]
    except Exception:
        return []
    out = []
    for r in rows:
        out.append({
            "t": r.get("Timestamp"),
            "users": _f(r.get("User Count")),
            "rps": _f(r.get("Requests/s")),
            "fps": _f(r.get("Failures/s")),
            "p50": _f(r.get("50%")),
            "p95": _f(r.get("95%")),
        })
    return out


def _f(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def build_reports(analysis: dict, run_dir: Path) -> dict:
    run_dir = Path(run_dir)
    reports = run_dir / "reports"
    reports.mkdir(parents=True, exist_ok=True)
    history = _read_history(run_dir)
    html_path = reports / "performance_report.html"
    xlsx_path = reports / "performance_report.xlsx"
    # FAIL-SAFE: a report must ALWAYS be produced. If full rendering throws (an
    # edge in the run's data — e.g. an aborted / 0-order / payment-failed run), fall
    # back to a minimal but valid report that still carries the key metrics AND the
    # traceback, so a run is never left with no report and the cause is visible.
    try:
        html = _html(analysis, history)
    except Exception as exc:
        import traceback
        html = _fallback_html(analysis, exc, traceback.format_exc())
    html_path.write_text(html, encoding="utf-8")
    try:
        _xlsx(analysis, history, xlsx_path)
        xlsx_ok = True
    except Exception:
        xlsx_ok = False
    return {"html": str(html_path), "xlsx": str(xlsx_path) if xlsx_ok else None}


def _fallback_html(analysis: dict, exc: Exception, tb: str) -> str:
    """Minimal, always-valid report used when full rendering fails, so the run is
    never left with no report. Shows headline metrics + the error/traceback."""
    a = analysis or {}
    o = a.get("overall") or {}
    sla = a.get("sla") or {}
    flow = a.get("flow") or {}
    bt = (a.get("browser_track") or {}).get("summary") or {}
    def _row(k, v):
        return "<tr><td>%s</td><td>%s</td></tr>" % (html.escape(str(k)), html.escape(str(v)))
    rows = "".join([
        _row("Total requests", o.get("total_requests", 0)),
        _row("Failures", o.get("total_failures", 0)),
        _row("Error rate %", o.get("error_rate", 0)),
        _row("p95 (ms)", o.get("p95", 0)),
        _row("Throughput (req/s)", o.get("throughput", 0)),
        _row("SLA pass", sla.get("pass", "-")),
        _row("Orders created", flow.get("orders", "-")),
        _row("Login ok / fail", "%s / %s" % (flow.get("login_ok", "-"), flow.get("login_fail", "-"))),
        _row("Checkout stop reason", (flow.get("checkout_state") or {}).get("stop_reason", "-")),
        _row("Browser track", "%s iter, %s fail — %s" % (
            bt.get("iterations", "-"), bt.get("failures", "-"), bt.get("note", ""))),
    ])
    return (
        "<!doctype html><html><head><meta charset='utf-8'>"
        "<title>LT Metrics Performance Report (fallback)</title>"
        "<style>body{font-family:Aptos,Inter,'Segoe UI',system-ui,sans-serif;"
        "margin:32px;background:#0A0A0A;color:#F4F4F4;line-height:1.5}"
        "table{border-collapse:collapse;margin:16px 0}td{border:1px solid #2A2A2A;"
        "padding:9px 12px}h1{font-size:30px;font-weight:300;margin:0 0 6px}"
        "pre{background:#000;border:1px solid #2A2A2A;padding:12px;overflow:auto;"
        "color:#EDEDED;font-size:12px;font-family:'JetBrains Mono',Consolas,monospace}"
        "</style></head>"
        "<body><h1>LT Metrics Performance Report</h1>"
        "<p style='color:#ABABAB'>The full report couldn't be rendered for this run "
        "(likely an aborted / 0-order / payment-failed edge). Key results below; the "
        "error is included so it can be fixed.</p>"
        "<table>" + rows + "</table>"
        "<h2 style='color:var(--accent)'>Rendering error</h2>"
        "<p>%s</p><pre>%s</pre>"
        "<p style='color:#ABABAB'>All raw results (locust_stats.csv, ltm_flow.json, "
        "browser_track.json, ltm_calls*.jsonl) are intact in this run's results folder.</p>"
        "</body></html>" % (html.escape(str(exc)), html.escape(tb)))


# --------------------------------------------------------------------------- #
# HTML
# --------------------------------------------------------------------------- #
# A tile's value is plain text unless the caller has something to say about
# it. The default used to be the brand red, so every figure on the page read
# as an alarm -- including "Total Requests".
def _kpi(label, value, sub="", accent="var(--txt)"):
    return f'''<div class="kpi"><div class="kpi-label">{label}</div>
      <div class="kpi-value" style="color:{accent}">{value}</div>
      <div class="kpi-sub">{sub}</div></div>'''


def _html(a: dict, history: list[dict]) -> str:
    o = a["overall"]
    sla = a["sla"]
    plan = a["plan"]
    gate_pass = sla["pass"]
    gate_color = "var(--ok)" if gate_pass else "var(--sev3)"
    gate_text = "PASS ✅" if gate_pass else "FAIL ❌"

    kpis = "".join([
        _kpi("Total Requests", f"{o['total_requests']:,}"),
        _kpi("Failures", f"{o['total_failures']:,}",
             f"{o['error_rate']}%", "var(--sev3)" if o["error_rate"] > 0 else "var(--ok)"),
        _kpi("Error Rate", f"{o['error_rate']}%", f"budget {sla['max_error_rate_pct']}%",
             "var(--ok)" if sla["error_gate"] else "var(--sev3)"),
        _kpi("Throughput", f"{o['throughput']}", "req/s"),
        _kpi("Avg Response", f"{o['avg_response']} ms"),
        _kpi("p95 Response", f"{o['p95']} ms", f"target {int(sla['max_p95_ms'])} ms",
             "var(--ok)" if sla["p95_gate"] else "var(--sev3)"),
        _kpi("p99 Response", f"{o['p99']} ms"),
        _kpi("Max Response", f"{o['max_response']} ms"),
    ])

    # checkout-flow KPIs (orders created / login) when a recording flow ran
    flow = a.get("flow") or {}
    flow_html = ""
    if flow:
        orders = flow.get("orders", 0)
        lok = flow.get("login_ok", 0)
        lfail = flow.get("login_fail", 0)
        heals = flow.get("heals", 0)
        captcha = flow.get("captcha", 0)
        # Order value captured at checkout (base_grand_total) — real revenue, not £0.
        order_value_total = flow.get("order_value_total")
        _order_values = flow.get("order_values") or []
        _avg_order = round(order_value_total / len(_order_values), 2) if _order_values else 0
        flow_html = ("<h2>Checkout Flow</h2><div class=\"kpis\">"
                     + _kpi("Orders Created", orders, "confirmed order ids",
                            "var(--ok)" if orders else "var(--sev1)")
                     + (_kpi("Order Value (total)", order_value_total,
                             "sum of base_grand_total across orders", "var(--ok)")
                        + _kpi("Avg Order Value", _avg_order,
                               "mean order value at checkout", "var(--sev0)")
                        if order_value_total is not None else "")
                     + _kpi("Login Success", lok, "valid credentials", "var(--ok)")
                     + _kpi("Login Failed", lfail, "rejected logins",
                            "var(--accent)" if lfail else "var(--mut)")
                     + _kpi("Auto-heals", heals, "run-time self-corrections",
                            "var(--sev0)" if heals else "var(--mut)")
                     + _kpi("CAPTCHA hits", captcha, "challenges blocking the flow",
                            "var(--accent)" if captcha else "var(--mut)")
                     + "</div>")
        # Effective run configuration — the profile the script ACTUALLY ran with,
        # so the report is self-describing and comparable across runs. Highlights
        # STRICT (reproducible) mode when it was on.
        _eff = flow.get("effective_profile") or {}
        if _eff:
            _strict_on = bool(_eff.get("strict"))
            _badge = ("<b style=\"color:var(--ok)\">STRICT / reproducible</b>" if _strict_on
                      else "<b style=\"color:var(--mut)\">adaptive (default)</b>")
            _parts = [
                "users=%s" % html.escape(str(_eff.get("users", "-"))),
                "duration=%s" % html.escape(str(_eff.get("duration", "-"))),
                "think=%s-%ss" % (html.escape(str(_eff.get("think_min", "-"))),
                                  html.escape(str(_eff.get("think_max", "-")))),
                "payment=%s" % html.escape(str(_eff.get("forced_payment") or "auto")),
                "data_sharing=%s" % html.escape(str(_eff.get("data_sharing", "-"))),
                "faithful=%s" % html.escape(str(_eff.get("faithful", "-"))),
                "build=%s" % html.escape(str(_eff.get("build", "-"))),
            ]
            flow_html += ("<p class=\"muted\" style=\"margin-top:8px\">"
                          "<b>Run configuration (effective):</b> " + _badge + " &middot; "
                          + " &middot; ".join(_parts) + "</p>")
        oids = flow.get("order_ids") or []
        if oids:
            flow_html += ("<p class=\"muted\" style=\"margin-top:8px\"><b>Order ids created:</b> "
                          + html.escape(", ".join(str(x) for x in oids[:60])) + "</p>")
        if captcha:
            flow_html += ("<p class=\"muted\" style=\"margin-top:8px;color:var(--accent)\">"
                          "<b>CAPTCHA blocked the flow.</b> A CAPTCHA cannot be solved by a "
                          "load test. Disable it on the test environment, use the provider's "
                          "test keys, allowlist the load-generator IPs, or supply a bypass "
                          "token in Advanced overrides.</p>")

    # endpoint table (+ note on any static-asset rows excluded from the breakdown)
    _se = a.get("static_excluded") or {}
    static_note = ""
    if _se.get("count"):
        static_note = ('<p style="font-size:12px;color:var(--mut);margin-top:8px">'
                       '%d static-asset endpoint(s) (%s requests) — html / css / js / images / '
                       'fonts — excluded from this API breakdown.</p>'
                       % (_se["count"], format(int(_se.get("requests", 0)), ",")))
    # Track B (Playwright) section — only when a browser track actually ran.
    browser_html = ""
    _bt = a.get("browser_track")
    if _bt:
        _s = _bt.get("summary") or {}
        _brows = ""
        for e in _bt.get("endpoints", []):
            _brows += ("<tr><td>%s</td><td class=\"num\">%s</td><td class=\"num\">%s</td>"
                       "<td class=\"num\">%d</td><td class=\"num\">%d</td></tr>" % (
                        html.escape(str(e.get("name", ""))),
                        format(int(e.get("num_requests", 0)), ","),
                        format(int(e.get("num_failures", 0)), ","),
                        int(e.get("avg", 0)), int(e.get("p95", 0))))
        _tbl = ("<table><thead><tr><th>Transaction</th><th>Samples</th><th>Fails</th>"
                "<th>Avg</th><th>p95</th></tr></thead><tbody>%s</tbody></table>" % _brows) if _brows else ""
        # real (client-rendered) prices the browser saw — the ones HTTP replay reports as 0
        _prices = _s.get("prices") or []
        _price_html = ""
        if _prices:
            _seen, _items = set(), []
            for p in _prices:
                v = str(p.get("price", ""))
                if v and v not in _seen:
                    _seen.add(v)
                    _items.append(html.escape(v))
            if _items:
                _price_html = ("<p style=\"color:var(--mut);font-size:13px\">Real product "
                               "price(s) captured in-browser (HTTP replay sees 0): <b>%s</b></p>"
                               % ", ".join(_items[:10]))
        # Crawl-derived order-confirmation assertion: the ground-truth success
        # signal the real browser observed. This is the authoritative "order placed"
        # assertion; the HTTP load script should confirm orders the same way.
        _oc = _s.get("order_confirmation") or {}
        _oc_html = ""
        if _oc.get("confirmed"):
            _bits = []
            if _oc.get("signal_type") and _oc.get("signal"):
                _bits.append("%s <code>%s</code>" % (html.escape(str(_oc["signal_type"])),
                                                     html.escape(str(_oc["signal"]))))
            if _oc.get("order_number"):
                _bits.append("order # <b>%s</b>" % html.escape(str(_oc["order_number"])))
            if _oc.get("url"):
                _bits.append(html.escape(str(_oc["url"])))
            _oc_html = ("<p style=\"color:var(--ok);font-size:13px\">Order-confirmation "
                        "assertion observed by the browser: %s. Ensure this signal is in "
                        "<code>platform_rules.order_url_signals</code> so the HTTP load "
                        "script asserts order completion the same way.</p>"
                        % ("; ".join(_bits) or "confirmed"))
        browser_html = (
            "<h2>Browser Track (Playwright)</h2>"
            "<p style=\"color:var(--mut);font-size:13px\">Real-browser track — "
            "%s iteration(s), %s request(s), %s failure(s). %s</p>%s%s%s" % (
                format(int(_s.get("iterations", 0)), ","),
                format(int(_s.get("requests", 0)), ","),
                format(int(_s.get("failures", 0)), ","),
                html.escape(str(_s.get("note", ""))), _price_html, _oc_html, _tbl))

    # Business Data Dependencies — Browser-to-API fidelity model.
    business_html = ""
    _bd = a.get("business_data")
    if _bd and _bd.get("dependencies"):
        _cls_color = {"PARAMETER": "var(--sev0)", "CORRELATION": "var(--ok)",
                      "RUNTIME_DERIVED": "var(--sev1)", "CLIENT_CALCULATED": "var(--sev2)",
                      "SERVER_GENERATED": "var(--mut)", "BUSINESS_REFERENCE": "var(--txt)",
                      "STATIC": "var(--mut)", "UNKNOWN": "var(--accent)"}
        _rows = ""
        for d in _bd["dependencies"]:
            cls = str(d.get("classification", "UNKNOWN"))
            src = d.get("source") or {}
            origin = (src.get("origin") or {}).get("type") or src.get("resolved_from") or "-"
            api_v = d.get("api_value"); br_v = d.get("browser_value")
            val = d.get("value")
            shown = (("API %s / Browser %s" % (api_v, br_v)) if (api_v is not None or br_v is not None)
                     else (str(val) if val is not None else "-"))
            flag = " ⚠️ mismatch" if d.get("mismatch") else ""
            _rows += ("<tr><td>%s</td><td><span style=\"color:%s;font-weight:700\">%s</span></td>"
                      "<td>%s</td><td>%s%s</td></tr>" % (
                          html.escape(str(d.get("name", ""))),
                          _cls_color.get(cls, "var(--txt)"), html.escape(cls),
                          html.escape(str(origin)), html.escape(str(shown)), flag))
        _gate = _bd.get("gate") or {}
        _gate_html = ("" if _gate.get("ok", True) else
                      "<p style=\"color:var(--sev2);font-size:13px\"><b>Fidelity gate:</b> %s</p>"
                      % html.escape(str(_gate.get("reason", ""))))
        business_html = (
            "<h2>Business Data Dependencies</h2>"
            "<p style=\"color:var(--mut);font-size:13px\">How each business value is "
            "sourced to faithfully reproduce the transaction at API level, and where "
            "browser and API disagree.</p>%s"
            "<table><thead><tr><th>Value</th><th>Classification</th><th>Source</th>"
            "<th>Observed</th></tr></thead><tbody>%s</tbody></table>" % (_gate_html, _rows))

    # Payment Replay Analysis — can the card step run at the HTTP layer? Built from
    # the captured payment network sequence (Track B) + KB replay profiles.
    payment_replay_html = ""
    _pr = a.get("payment_replay")
    if _pr:
        _rv = _pr.get("replayable")
        _rv_color = {True: "var(--ok)", "true": "var(--ok)", "conditional": "var(--sev1)",
                     False: "var(--accent)", "false": "var(--accent)"}.get(_rv, "var(--mut)")
        _rv_label = {True: "REPLAYABLE", "true": "REPLAYABLE", "conditional": "CONDITIONAL",
                     False: "NOT REPLAYABLE", "false": "NOT REPLAYABLE"}.get(_rv, "UNDETERMINED")
        _cap_rows = ""
        for c in (_pr.get("captured") or [])[:15]:
            _cap_rows += ("<tr><td>%s</td><td style=\"font-size:11px\">%s</td>"
                          "<td class=\"num\">%s</td></tr>" % (
                              html.escape(str(c.get("method", ""))),
                              html.escape(str(c.get("url", ""))[:120]),
                              html.escape(str(c.get("status") if c.get("status") is not None else "-"))))
        _cap_tbl = ("<table><thead><tr><th>Method</th><th>Endpoint (redacted)</th>"
                    "<th>Status</th></tr></thead><tbody>%s</tbody></table>" % _cap_rows) if _cap_rows else (
                    "<p style=\"color:var(--mut);font-size:13px\">No payment traffic was "
                    "captured (Track B may not have reached the card step).</p>")
        def _bullets(items):
            items = [html.escape(str(x)) for x in (items or [])]
            return ("<ul style=\"font-size:13px;color:var(--mut)\">%s</ul>"
                    % "".join("<li>%s</li>" % i for i in items)) if items else ""
        _mint = _pr.get("mint") or {}
        _mint_html = ""
        if _mint:
            _md = _mint.get("description") or ""
            _me = _mint.get("resolved_mint_endpoint") or _mint.get("mint_endpoint_hint") or ""
            _mint_html = ("<p style=\"font-size:13px\"><b>Per-VU mint recipe:</b> %s%s</p>" % (
                html.escape(str(_md)),
                (" <code>%s</code>" % html.escape(str(_me))) if _me else ""))
        payment_replay_html = (
            "<h2>Payment Replay Analysis</h2>"
            "<p style=\"font-size:14px\"><span style=\"color:%s;font-weight:800\">%s</span> "
            "&nbsp;<span style=\"color:var(--mut);font-size:12px\">(%s confidence)</span></p>"
            "<p style=\"color:var(--mut);font-size:13px\">%s</p>"
            "%s%s"
            "<p style=\"font-size:13px;margin-top:10px\"><b>Why:</b></p>%s"
            "%s"
            "<p style=\"color:var(--mut);font-size:12px\">Card data and secrets are "
            "redacted at capture; public sandbox test cards only.</p>" % (
                _rv_color, _rv_label, html.escape(str(_pr.get("confidence", "medium"))),
                html.escape(str(_pr.get("verdict", ""))),
                _mint_html, _cap_tbl,
                _bullets(_pr.get("reasons")),
                (("<p style=\"font-size:13px;margin-top:6px\"><b>Blockers:</b></p>%s"
                  % _bullets(_pr.get("blockers"))) if _pr.get("blockers") else "")))

    # Payment Profile — payment as a first-class discovered dependency.
    payment_profile_html = ""
    _pp = a.get("payment_profile")
    if _pp:
        _strat_color = {"api_replay": "var(--ok)", "browser_assisted": "var(--sev1)",
                        "hybrid": "var(--sev2)", "offline": "var(--mut)"}.get(_pp.get("strategy"), "var(--mut)")
        _corr = _pp.get("correlation") or {}
        _tok = _pp.get("token") or {}
        _corr_line = ""
        if _corr:
            _prov = "proven" if _corr.get("proven") else "heuristic (%s confidence)" % _corr.get("confidence", "medium")
            _corr_line = ("<p style=\"font-size:13px\"><b>Token correlation:</b> "
                          "<code>%s</code> from <code>%s</code> &rarr; <code>%s</code> in "
                          "<code>%s</code> &nbsp;<span style=\"color:var(--mut)\">(%s)</span></p>" % (
                              html.escape(str(_corr.get("extract") or _corr.get("token_field") or "token")),
                              html.escape(str(_corr.get("token_endpoint") or "provider")),
                              html.escape(str(_corr.get("inject") or "?")),
                              html.escape(str(_corr.get("consumer_endpoint") or "?")),
                              html.escape(_prov)))
        _rowp = ("<tr><td>Provider</td><td>%s</td></tr>"
                 "<tr><td>Integration</td><td>%s</td></tr>"
                 "<tr><td>Token endpoint</td><td>%s</td></tr>"
                 "<tr><td>Replayable</td><td>%s</td></tr>"
                 "<tr><td>Strategy</td><td><span style=\"color:%s;font-weight:700\">%s</span></td></tr>" % (
                     html.escape(str(_pp.get("provider", "-"))),
                     html.escape(str(_pp.get("integration_type", "-"))),
                     html.escape(str(_tok.get("endpoint") or "-")),
                     html.escape(str(_pp.get("replayable", "-"))),
                     _strat_color, html.escape(str(_pp.get("strategy", "-")))))
        payment_profile_html = (
            "<h2>Payment Profile</h2>"
            "<p style=\"color:var(--mut);font-size:13px\">Payment treated as a "
            "first-class dependency — how the token is produced and consumed, and the "
            "execution strategy that follows.</p>"
            "<table><tbody>%s</tbody></table>%s"
            "<p style=\"color:var(--mut);font-size:12px\">%s</p>" % (
                _rowp, _corr_line, html.escape(str(_pp.get("summary", "")))))

    rows = ""
    for e in a["endpoints"]:
        # sla_pass is None for a transaction timer -- no target, so no verdict.
        # Rendering that as a cross blamed the journey for being a journey.
        _tgt = e.get("sla_target")
        status = "—" if _tgt is None else ("✅" if e.get("sla_pass") else "❌")
        err = (e["num_failures"] / e["num_requests"] * 100) if e["num_requests"] else 0
        rows += f'''<tr>
          <td>{html.escape(str(e["name"]))}</td>
          <td style="font-size:11px;color:var(--mut)">{html.escape(str(e.get("req_method","")))} {html.escape(str(e.get("endpoint","")))}</td>
          <td class="num">{e["num_requests"]:,}</td>
          <td class="num">{e["num_failures"]:,}</td>
          <td class="num">{err:.2f}%</td>
          <td class="num">{int(e["avg"])}</td>
          <td class="num">{int(e["p50"])}</td>
          <td class="num">{int(e["p90"])}</td>
          <td class="num">{int(e["p95"])}</td>
          <td class="num">{int(e["p99"])}</td>
          <td class="num">{e["rps"]:.2f}</td>
          <td class="num">{"—" if _tgt is None else int(_tgt)}</td>
          <td>{status}</td></tr>'''

    recs = ""
    for r in (a.get("recommendations") or []):
        # recommendations come from several sources (perf analyzer, KB, the
        # payment-gate) with DIFFERENT shapes — some carry impact/effort, some
        # only priority/title/detail. Render defensively so a missing key never
        # breaks the whole report.
        _prio = str(r.get("priority", "P3"))
        _impact = r.get("impact")
        _effort = r.get("effort")
        recs += (
            '<div class="rec">'
            '<span class="badge b-%s">%s</span> <b>%s</b>' % (
                html.escape(_prio), html.escape(_prio), html.escape(str(r.get("title", ""))))
            + ('<span class="tag">Impact: %s</span>' % html.escape(str(_impact)) if _impact else "")
            + ('<span class="tag">Effort: %s</span>' % html.escape(str(_effort)) if _effort else "")
            + '<div class="rec-detail">%s</div></div>' % html.escape(str(r.get("detail", ""))))

    breaches = ""
    if sla["breaches"]:
        for b in sla["breaches"]:
            breaches += f"<tr><td>{html.escape(b['name'])}</td><td class='num'>{int(b['p95'])}</td><td class='num'>{int(b['target'])}</td></tr>"
        breaches = f'''<table><thead><tr><th>Transaction</th><th>p95 (ms)</th>
          <th>Target (ms)</th></tr></thead><tbody>{breaches}</tbody></table>'''
    else:
        breaches = "<p class='ok'>No endpoint SLA breaches.</p>"

    fails = ""
    for f in a["failures"][:60]:
        # show the COMPLETE failure message (wrapped), not a truncated preview
        msg = html.escape(str(f['error'])).replace("\n", "<br>")
        fails += (f"<tr><td>{html.escape(str(f['name']))}</td>"
                  f"<td style='white-space:pre-wrap;word-break:break-word;"
                  f"font-family:var(--mono);font-size:11px'>{msg}</td>"
                  f"<td class='num'>{f['occurrences']}</td></tr>")
    fails = fails or "<tr><td colspan='3' class='ok'>No failures recorded.</td></tr>"

    trend = a["trend"]
    trend_html = f"<p>{html.escape(trend['note'])}</p>"
    if trend.get("has_baseline"):
        trend_html += (f"<p>p95 Δ {trend['p95_delta_pct']}% · error Δ "
                       f"{trend['error_delta_pct']}% · throughput Δ "
                       f"{trend['throughput_delta_pct']}% vs baseline "
                       f"{html.escape(str(trend['baseline_date']))}</p>")

    jira = ""
    if a.get("jira_markdown"):
        jira = f'''<h2>Jira Ticket (ready to copy)</h2>
          <pre class="jira">{html.escape(a["jira_markdown"])}</pre>'''

    _target = html.escape(str(a.get("target", "")))
    _domain = html.escape(str(a.get("domain", "")))
    _plabel = html.escape(plan["label"])
    _pusers = plan["users"]
    _pdur = plan["duration_human"]
    _rca_summary = html.escape(a["rca"]["summary"])
    _llm_txt = a.get("llm_rca")
    _llm_rca_html = (
        "<h2>AI Root-Cause Analysis (Claude)</h2><div class=\"rec\">"
        + html.escape(_llm_txt).replace("\n", "<br>") + "</div>") if _llm_txt else ""

    # Raw epoch seconds ("1788535425") tell a reader nothing, and the console
    # plots elapsed time -- so the same run had two different axes. Elapsed,
    # both places.
    def _elapsed(v, t0):
        try:
            return "%ds" % (int(float(v)) - t0)
        except (TypeError, ValueError):
            return str(v)
    try:
        _t0 = int(float(history[0]["t"])) if history else 0
    except (TypeError, ValueError, KeyError, IndexError):
        _t0 = 0
    labels = json.dumps([_elapsed(h.get("t"), _t0) for h in history])
    p50 = json.dumps([h["p50"] for h in history])
    p95 = json.dumps([h["p95"] for h in history])
    users = json.dumps([h["users"] for h in history])
    rps = json.dumps([h["rps"] for h in history])
    fps = json.dumps([h["fps"] for h in history])
    ep_labels = json.dumps([e["name"] for e in a["endpoints"]])
    ep_p95 = json.dumps([e["p95"] for e in a["endpoints"]])
    ep_target = json.dumps([e.get("sla_target", 0) for e in a["endpoints"]])

    return f'''<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>LT Metrics Performance Report — {_target}</title>
<script src="https://cdn.jsdelivr.net/npm/chart.js@4"></script>
<link rel="preconnect" href="https://fonts.googleapis.com">
<!-- Aptos is the design system's face and ships with Office/Windows, so it
     is not fetched here; Inter is the system's own next choice and needs 300
     for the Light headings. Space Grotesk is gone -- display type shares the
     body stack and differs from it only by weight. -->
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap" rel="stylesheet">
<style>
  /* AIBP Design System V2.0 / canvas 9a -- the same tokens the console uses,
     so a report reads as the same product and not a cousin of it. */
  :root{{
    --bg:#0A0A0A;--panel:#131313;--panel2:#1C1C1C;--line:#767676;--hair:#2A2A2A;
    --txt:#F4F4F4;--mut:#ABABAB;--accent:#C91A12;--green:#23674A;--orange:#EF9253;
    --ok:#5CD68F;--eyebrow:#EF9253;--mcf-red:#E00000;--logbg:#000000;
    --sev0:#8E93B4;--sev1:#C6BF52;--sev2:#E0834A;--sev3:#FF7B70;--sev4:#C0392B;
    --head:Aptos,Inter,'Segoe UI',system-ui,-apple-system,Helvetica,Arial,sans-serif;
    --mono:'JetBrains Mono','SF Mono',Consolas,'Liberation Mono',monospace;
    --ls-overline:.20em;--ls-caps:.08em;--fw-light:300;--fw-bold:700;
  }}
  *{{box-sizing:border-box}}
  body{{margin:0;background:var(--bg);color:var(--txt);font-family:var(--head);
    font-size:16px;line-height:1.5;border-top:3px solid var(--accent)}}
  .masthead{{display:flex;align-items:baseline;gap:14px;padding:14px 28px;
    background:var(--green);color:#fff;flex-wrap:wrap}}
  .masthead .mk{{font-family:var(--mono);font-size:17px;letter-spacing:.1em;
    font-weight:var(--fw-bold);line-height:1}}
  .masthead .ms{{font-size:12px;font-weight:var(--fw-bold);text-transform:uppercase;
    letter-spacing:.16em;line-height:1;color:var(--orange)}}
  .wrap{{padding:28px}}
  /* Display headings Light, structural ones Bold -- the system's own split. */
  h1{{font-size:30px;font-weight:var(--fw-light);line-height:1.2;
    letter-spacing:-.02em;margin:0 0 6px;max-width:640px}}
  h3{{font-size:20px;font-weight:var(--fw-bold);line-height:1.15;margin:0 0 6px}}
  /* A section heading is the canvas's eyebrow: caps, tracked, and led by the
     node-chain mark. It replaces the underlined h2 this page used. */
  h2{{font-size:12px;font-weight:var(--fw-bold);text-transform:uppercase;
    letter-spacing:var(--ls-overline);color:var(--eyebrow);
    display:flex;align-items:center;gap:14px;margin:32px 0 10px}}
  h2::before{{content:"";display:block;flex:none;width:58px;height:12px;
    background:
      radial-gradient(circle, var(--mcf-red) 5px, transparent 5px) 0 0/12px 12px no-repeat,
      linear-gradient(var(--mcf-red),var(--mcf-red)) 12px 5px/6px 2px no-repeat,
      radial-gradient(circle, var(--mcf-red) 5px, transparent 5px) 18px 0/12px 12px no-repeat,
      linear-gradient(var(--mcf-red),var(--mcf-red)) 30px 5px/6px 2px no-repeat,
      radial-gradient(circle, rgba(224,0,0,.45) 5px, transparent 5px) 36px 0/12px 12px no-repeat,
      linear-gradient(rgba(224,0,0,.6),rgba(224,0,0,.6)) 48px 5px/6px 2px no-repeat,
      radial-gradient(circle, var(--txt) 5px, transparent 5px) 54px 0/12px 12px no-repeat;
  }}
  a{{color:var(--txt);text-decoration:underline;text-underline-offset:2px}}
  a:hover{{color:var(--accent)}}
  .sub{{color:var(--mut);margin-bottom:14px;font-size:14px}}
  /* Square and 44px, like every control in the console. */
  .gate{{display:inline-flex;align-items:center;min-height:44px;padding:0 18px;
    font-family:var(--mono);font-weight:var(--fw-bold);font-size:14px;
    text-transform:uppercase;letter-spacing:var(--ls-caps);
    background:{gate_color};color:#191919}}
  /* Four-up, separated by a 1px rule rather than a gap. */
  .kpis{{display:grid;grid-template-columns:repeat(4,1fr);gap:1px;
    background:var(--hair);border:1px solid var(--hair);margin:20px 0}}
  .kpi{{background:var(--panel2);padding:15px}}
  .kpi-label{{color:var(--mut);font-size:11px;text-transform:uppercase;
    letter-spacing:var(--ls-caps)}}
  .kpi-value{{font-size:23px;font-weight:var(--fw-bold);margin:4px 0 0}}
  .kpi-sub{{color:var(--mut);font-size:11px;margin-top:4px}}
  table{{width:100%;border-collapse:collapse;background:var(--panel);
    font-size:13px;border:1px solid var(--hair)}}
  th,td{{padding:9px 12px;text-align:left;border-bottom:1px solid var(--hair)}}
  th{{background:var(--panel2);color:var(--mut);font-weight:600;font-size:11px;
    text-transform:uppercase;letter-spacing:var(--ls-caps)}}
  /* A right-aligned column needs a right-aligned heading, or every word sits a
     column away from its own figures. */
  th.num,td.num{{text-align:right}}
  td.num{{font-variant-numeric:tabular-nums;font-family:var(--mono)}}
  tr:last-child td{{border-bottom:none}}
  .grid2{{display:grid;grid-template-columns:1fr 1fr;gap:20px}}
  .chart-card{{background:var(--panel2);border:1px solid var(--hair);padding:18px}}
  .rec{{background:var(--panel2);border:1px solid var(--hair);
    border-left:4px solid var(--accent);padding:12px 14px;margin:10px 0;
    font-size:14px}}
  .rec-detail,.d{{color:var(--mut);margin-top:6px;font-size:13px}}
  /* Priorities ride the severity ladder instead of five unrelated hues. */
  .badge{{font-family:var(--mono);font-weight:var(--fw-bold);padding:3px 9px;
    margin-right:8px;font-size:11px;color:#191919}}
  .b-P1{{background:var(--sev3)}} .b-P2{{background:var(--sev2)}}
  .b-P3{{background:var(--sev1)}} .b-P4{{background:var(--ok)}}
  .tag{{background:var(--panel);border:1px solid var(--hair);color:var(--mut);
    padding:2px 8px;font-size:12px;margin-right:6px;font-family:var(--mono)}}
  .ok{{color:var(--ok)}}
  .jira{{background:var(--logbg);border:1px solid var(--hair);padding:14px;
    white-space:pre-wrap;color:#EDEDED;font-family:var(--mono);font-size:12px}}
  footer{{margin:36px 28px 0;color:var(--mut);font-size:12px;text-align:center;
    border-top:1px solid var(--hair);padding:14px 0 28px}}
  @media(max-width:900px){{.kpis{{grid-template-columns:repeat(2,1fr)}}}}
  @media(max-width:520px){{.kpis{{grid-template-columns:1fr}}}}
  @media(max-width:800px){{.grid2{{grid-template-columns:1fr}}}}
  /* A report gets printed and mailed on. On paper the console's dark ground
     wastes ink and reads badly, so print takes the light half of the same
     palette rather than a different design. */
  @media print{{
    :root{{--bg:#FFFFFF;--panel:#FBEEE3;--panel2:#F4EDE5;--line:#B49A82;
      --hair:#E2D3C4;--txt:#191919;--mut:#5A4A3F;--eyebrow:#B34E14;
      --ok:#157347;--logbg:#191919;
      --sev0:#5C607A;--sev1:#8E8931;--sev2:#B4551F;--sev3:#C91A12;--sev4:#8A120C;}}
    body{{border-top:none}}
    .masthead{{color:#fff;-webkit-print-color-adjust:exact;print-color-adjust:exact}}
    h2{{break-after:avoid}}
    table,.chart-card,.rec,.kpis{{break-inside:avoid}}
  }}
</style></head>
<body>
  <div class="masthead"><span class="mk">LT METRICS</span>
    <span class="ms">Load Testing Reimagined</span></div>
  <div class="wrap">
  <h1>LT Metrics Performance Report</h1>
  <div class="sub">{_target} · {_domain}
    · {_plabel} · {_pusers} users · {_pdur}</div>
  <div class="gate">CI Quality Gate: {gate_text}</div>
  <div class="kpis">{kpis}</div>
  {flow_html}

  <h2>Trends</h2>
  <div class="grid2">
    <div class="chart-card"><canvas id="latency"></canvas></div>
    <div class="chart-card"><canvas id="load"></canvas></div>
  </div>
  <div class="chart-card" style="margin-top:20px"><canvas id="epchart"></canvas></div>

  <h2>Endpoint Validation & Statistics</h2>
  <table><thead><tr>
    <th>Transaction</th><th>Endpoint</th><th>Samples</th><th>Fails</th><th>Err %</th><th>Avg</th>
    <th>p50</th><th>p90</th><th>p95</th><th>p99</th><th>req/s</th><th>SLA (ms)</th>
    <th>Status</th></tr></thead><tbody>{rows}</tbody></table>
  <p class="sub" style="margin-top:8px">A row with no target is a transaction
  timer: it wraps the calls inside it, so its time is the sum of theirs and a
  per-request target does not apply to it. Its failures still count.</p>
  {static_note}
  {browser_html}
  {business_html}
  {payment_profile_html}
  {payment_replay_html}

  <h2>SLA Compliance</h2>
  {breaches}

  <h2>AI Recommendations</h2>
  {recs}

  <h2>Root Cause Analysis</h2>
  <div class="rec"><b>{_rca_summary}</b></div>

  {_llm_rca_html}

  <h2>Historical Trend</h2>
  <div class="rec">{trend_html}</div>

  <h2>Failures</h2>
  <table><thead><tr><th>Transaction</th><th>Error</th><th>Count</th></tr></thead>
    <tbody>{fails}</tbody></table>

  {jira}

  </div>
  <footer>Generated by LT Metrics — Load Testing Reimagined</footer>

<script>
const mk=(id,cfg)=>{{const el=document.getElementById(id);if(el)new Chart(el,cfg);}};
// Chart.js draws to a canvas, which cannot read a CSS custom property, so the
// tokens are resolved once here. Same assignment as the console: sev0 for
// volume, sev3 for latency, muted for users, a dashed line for a target.
const T=(n,f)=>(getComputedStyle(document.documentElement).getPropertyValue(n)||'').trim()||f;
const C={{txt:T('--txt','#F4F4F4'),mut:T('--mut','#ABABAB'),hair:T('--hair','#2A2A2A'),
  sev0:T('--sev0','#8E93B4'),sev1:T('--sev1','#C6BF52'),sev2:T('--sev2','#E0834A'),
  sev3:T('--sev3','#FF7B70'),accent:T('--accent','#C91A12')}};
const gopt=t=>({{responsive:true,plugins:{{title:{{display:true,text:t,color:C.txt}},legend:{{labels:{{color:C.mut}}}}}},scales:{{x:{{ticks:{{color:C.mut,maxTicksLimit:8}},grid:{{color:C.hair}}}},y:{{ticks:{{color:C.mut}},grid:{{color:C.hair}}}}}}}});
mk('latency',{{type:'line',data:{{labels:{labels},datasets:[
  {{label:'p50 (ms)',data:{p50},borderColor:C.sev1,tension:.3,pointRadius:0}},
  {{label:'p95 (ms)',data:{p95},borderColor:C.sev3,borderWidth:2.5,tension:.3,pointRadius:0}}]}},
  options:gopt('Response Time Over Time')}});
mk('load',{{type:'line',data:{{labels:{labels},datasets:[
  {{label:'Users',data:{users},borderColor:C.mut,borderDash:[6,4],tension:.3,pointRadius:0}},
  {{label:'Requests/s',data:{rps},borderColor:C.sev0,tension:.3,pointRadius:0}},
  {{label:'Failures/s',data:{fps},borderColor:C.sev3,tension:.3,pointRadius:0}}]}},
  options:gopt('Load & Throughput Over Time')}});
mk('epchart',{{type:'bar',data:{{labels:{ep_labels},datasets:[
  {{label:'p95 (ms)',data:{ep_p95},backgroundColor:C.sev0}},
  {{label:'SLA target (ms)',data:{ep_target},type:'line',borderColor:C.sev3,
    borderDash:[4,4],borderWidth:1,pointRadius:0}}]}},
  options:gopt('p95 per Transaction vs SLA')}});
</script>
</body></html>'''


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _xlsx(a: dict, history: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import LineChart, Reference

    wb = Workbook()
    o = a["overall"]
    sla = a["sla"]
    hdr = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1E2B42")
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "LT Metrics Performance Report"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Target", a.get("target")),
        ("Domain", a.get("domain")),
        ("Profile", a["plan"]["label"]),
        ("Users", a["plan"]["users"]),
        ("Duration", a["plan"]["duration_human"]),
        ("", ""),
        ("Total Requests", o["total_requests"]),
        ("Total Failures", o["total_failures"]),
        ("Error Rate %", o["error_rate"]),
        ("Avg Response (ms)", o["avg_response"]),
        ("p95 (ms)", o["p95"]),
        ("p99 (ms)", o["p99"]),
        ("Throughput (req/s)", o["throughput"]),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws[f"A{i}"] = k
        ws[f"B{i}"] = v
        ws[f"A{i}"].font = Font(bold=True)
    r = len(rows) + 4
    ws[f"A{r}"] = "CI Quality Gate"
    ws[f"A{r}"].font = Font(bold=True)
    ws[f"B{r}"] = "PASS" if sla["pass"] else "FAIL"
    ws[f"B{r}"].fill = green if sla["pass"] else red
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 40

    # Endpoint Stats
    ws2 = wb.create_sheet("Endpoint Stats")
    cols = ["Transaction", "Endpoint", "Samples", "Failures", "Err %", "Avg", "p50",
            "p90", "p95", "p99", "req/s", "SLA Target", "Status"]
    ws2.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws2.cell(row=1, column=c)
        cell.font = hdr
        cell.fill = hfill
    for e in a["endpoints"]:
        err = (e["num_failures"] / e["num_requests"] * 100) if e["num_requests"] else 0
        endpoint = f'{e.get("req_method", "")} {e.get("endpoint", "")}'.strip()
        ws2.append([e["name"], endpoint, e["num_requests"], e["num_failures"], round(err, 2),
                    round(e["avg"], 1), round(e["p50"], 1), round(e["p90"], 1),
                    round(e["p95"], 1), round(e["p99"], 1), round(e["rps"], 2),
                    ("" if e.get("sla_target") is None else e.get("sla_target")),
                    ("n/a" if e.get("sla_target") is None
                     else ("PASS" if e.get("sla_pass") else "FAIL"))])
        last = ws2.cell(row=ws2.max_row, column=len(cols))
        if e.get("sla_target") is not None:
            last.fill = green if e.get("sla_pass") else red
    for col in "ABCDEFGHIJKLM":
        ws2.column_dimensions[col].width = 14
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 40

    # SLA Compliance
    ws3 = wb.create_sheet("SLA Compliance")
    ws3.append(["Metric", "Observed", "Target", "Result"])
    for c in range(1, 5):
        ws3.cell(row=1, column=c).font = hdr
        ws3.cell(row=1, column=c).fill = hfill
    ws3.append(["Error Rate %", o["error_rate"], sla["max_error_rate_pct"],
                "PASS" if sla["error_gate"] else "FAIL"])
    ws3.append(["p95 (ms)", o["p95"], sla["max_p95_ms"],
                "PASS" if sla["p95_gate"] else "FAIL"])
    for row in ws3.iter_rows(min_row=2):
        row[3].fill = green if row[3].value == "PASS" else red
    for col in "ABCD":
        ws3.column_dimensions[col].width = 18

    # Recommendations
    ws4 = wb.create_sheet("Recommendations")
    ws4.append(["Priority", "Impact", "Effort", "Title", "Detail"])
    for c in range(1, 6):
        ws4.cell(row=1, column=c).font = hdr
        ws4.cell(row=1, column=c).fill = hfill
    for rec in a["recommendations"]:
        ws4.append([rec["priority"], rec["impact"], rec["effort"], rec["title"],
                    rec["detail"]])
    ws4.column_dimensions["D"].width = 40
    ws4.column_dimensions["E"].width = 70

    # Failures
    ws5 = wb.create_sheet("Failures")
    ws5.append(["Transaction", "Error", "Occurrences"])
    for c in range(1, 4):
        ws5.cell(row=1, column=c).font = hdr
        ws5.cell(row=1, column=c).fill = hfill
    for f in a["failures"]:
        ws5.append([f["name"], str(f["error"]), f["occurrences"]])
    ws5.column_dimensions["A"].width = 24
    ws5.column_dimensions["B"].width = 70

    # Time Series + chart
    ws6 = wb.create_sheet("Time Series")
    ws6.append(["Timestamp", "Users", "Requests/s", "Failures/s", "p50", "p95"])
    for c in range(1, 7):
        ws6.cell(row=1, column=c).font = hdr
        ws6.cell(row=1, column=c).fill = hfill
    for h in history:
        ws6.append([h["t"], h["users"], h["rps"], h["fps"], h["p50"], h["p95"]])
    if len(history) > 1:
        chart = LineChart()
        chart.title = "Response Time Over Time"
        data = Reference(ws6, min_col=5, max_col=6, min_row=1, max_row=len(history) + 1)
        chart.add_data(data, titles_from_data=True)
        ws6.add_chart(chart, "H2")
    for col in "ABCDEF":
        ws6.column_dimensions[col].width = 16

    wb.save(path)
